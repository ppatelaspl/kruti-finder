"""Write the results workbook and merge approved new Kruti into an updated Kruti master.

One file, Kruti_Results.xlsx:
  New Kruti   - every Kruti found in the books that is not in the Excel; one row per book
                and page range, rows of the same Kruti kept together under one Temp ID
  Excel Kruti - the Excel's own Kruti: where each was found, or Not Found
"""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.worksheet.datavalidation import DataValidation

RESULTS_FILE = "Kruti_Results.xlsx"
NEW_SHEET, KNOWN_SHEET = "New Kruti", "Excel Kruti"

# Aspire brand palette only
GOLD, ORANGE, RED, PINK, PLUM, NAVY, GREY = (
    "F8BB13", "EE7532", "F04735", "D6175C", "B81A58", "011D45", "B2BBC7")
BAND = "EEF1F5"          # light shade marking every other Kruti group
FONT = "Outfit"

_HEADER_FILL = PatternFill("solid", fgColor=NAVY)
_INPUT_FILL = PatternFill("solid", fgColor=GREY)
_BAND_FILL = PatternFill("solid", fgColor=BAND)
_STATUS_STYLE = {
    "Found": (GOLD, NAVY),
    "Partially Found": (ORANGE, "FFFFFF"),
    "Not Found": (PINK, "FFFFFF"),
}
_THIN = Side(style="thin", color=GREY)
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP = Alignment(wrap_text=True, vertical="top")


def _table(ws, headers, rows, widths, wrap_cols=(), status_col=None, input_cols=(), group_col=None):
    """group_col: rows sharing a value in this column are one group; groups alternate shade."""
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = _HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = _BORDER
    for row in rows:
        ws.append([ILLEGAL_CHARACTERS_RE.sub(" ", v) if isinstance(v, str) else v for v in row])
    band, prev = False, object()
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if group_col:
            key = r[group_col - 1].value
            if key != prev:
                band, prev = not band, key
        for c in r:
            c.font = Font(name=FONT, color=NAVY)
            c.border = _BORDER
            c.alignment = _WRAP if c.column in wrap_cols else Alignment(vertical="top")
            if c.column in input_cols:
                c.fill = _INPUT_FILL
            elif group_col and band:
                c.fill = _BAND_FILL
        if status_col:
            cell = r[status_col - 1]
            if cell.value in _STATUS_STYLE:
                fill, color = _STATUS_STYLE[cell.value]
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = Font(name=FONT, bold=True, color=color)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = "1:1"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def _page(v):
    return "-" if v is None else v


# ------------------------------------------------------------------ results
NEW_COLUMNS = ["Temp ID"]            # + the Excel's 4 columns + these:
NEW_EXTRA = ["File Name", "Start Page", "End Page", "No. of Verses", "No. of Books",
             "OCR Quality", "Approve (Y/N)", "Remarks"]
KNOWN_EXTRA = ["Status", "File Name", "Start Page", "End Page", "Note"]


def write_results(path, groups, krutis, results, source_columns):
    wb = Workbook()
    cols = list(source_columns[:4])

    # ---- sheet 1: new Kruti, one row per book/page range, grouped by Temp ID
    ws = wb.active
    ws.title = NEW_SHEET
    rows = []
    for g in groups:
        copies = [g] + g.copies
        books = len({c.book for c in copies})
        for c in sorted(copies, key=lambda c: (c.book, c.start_page)):
            rows.append([g.temp_id, "", c.title, c.aadi, c.ant, c.book, c.start_page,
                         c.end_page, c.verses, books, c.ocr_quality, "", ""])
    n = len(NEW_COLUMNS) + len(cols)           # 5: Temp ID + Excel's 4 columns
    approve_col = n + NEW_EXTRA.index("Approve (Y/N)") + 1
    ocr_col = n + NEW_EXTRA.index("OCR Quality") + 1
    _table(ws, NEW_COLUMNS + cols + NEW_EXTRA, rows,
           [11, 11, 28, 50, 50, 42, 9, 9, 9, 9, 14, 11, 30],
           wrap_cols=(3, 4, 5, 6, n + len(NEW_EXTRA)), group_col=1,
           input_cols=(2, approve_col, approve_col + 1))
    if rows:
        dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
        ws.add_data_validation(dv)
        letter = ws.cell(1, approve_col).column_letter
        dv.add(f"{letter}2:{letter}{ws.max_row}")
    for r in ws.iter_rows(min_row=2):
        if r[ocr_col - 1].value != "OK":
            r[ocr_col - 1].font = Font(name=FONT, bold=True, color=RED)

    # ---- sheet 2: the Excel's Kruti, one row per book they were found in
    ws = wb.create_sheet(KNOWN_SHEET)
    rows = []
    for k in krutis:
        occ = results.get(k["id"], [])
        found = [o for o in occ if o.status == "Found"]
        if not occ:
            rows.append(list(k["raw"][:4]) + ["Not Found", "", "", "", "Not located in any book"])
        seen = set()
        for o in found + [o for o in occ if o.status != "Found"]:
            where = (o.book, o.start_page, o.end_page)
            if where in seen:            # the same book and pages once, even if matched twice
                continue
            seen.add(where)
            rows.append(list(k["raw"][:4]) + [o.status, o.book, _page(o.start_page),
                                              _page(o.end_page), o.note])
    _table(ws, cols + KNOWN_EXTRA, rows, [11, 28, 50, 50, 16, 42, 9, 9, 36],
           wrap_cols=(2, 3, 4, 6, 9), status_col=5)
    wb.save(path)


# ------------------------------------------------------------------ merge
def merge_approved(master_rows, headers, results_path, out_path):
    """Append each new Kruti approved (Y) in the results once, after the Excel's own rows."""
    ws = load_workbook(results_path)[NEW_SHEET]
    head = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    rows = [list(r) for r in master_rows]
    seen, added = set(), 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        temp_id = r[col["Temp ID"]]
        if str(r[col["Approve (Y/N)"]] or "").strip().upper() != "Y" or temp_id in seen:
            continue
        seen.add(temp_id)
        added += 1
        new = [r[1] or temp_id, r[2] or "", r[3], r[4]]      # Kruti No., name, Aadi, Ant
        rows.append((new + [""] * len(headers))[:len(headers)])

    out = Workbook()
    ws = out.active
    ws.title = "Kruti Master"
    _table(ws, list(headers), rows, [11, 30, 55, 55] + [16] * max(0, len(headers) - 4),
           wrap_cols=(2, 3, 4))
    out.save(out_path)
    return added, len(rows)
