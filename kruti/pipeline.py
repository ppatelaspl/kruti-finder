"""The full run - read the Excel, read the books one by one, split them into Kruti, flag
each Kruti as in the Excel (same or other script) or new, write one results Excel.

Used by both the command line and the desktop app. Progress is reported through a
callback; a Controller lets the caller pause, resume or cancel between pages.
"""
import os
import re
import threading
import time
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .extract import extract_book, page_count
from .krutis import Kruti, assign_groups, find_krutis
from .matcher import FOUND, Book, BookIndex, locate
from .normalize import devanagari_to_gujarati, script_of
from .report import RESULTS_FILE, write_results


class Cancelled(Exception):
    pass


class Controller:
    """Thread-safe pause / cancel switch checked between pages."""

    def __init__(self):
        self._run = threading.Event()
        self._run.set()
        self._cancel = threading.Event()

    def pause(self):
        self._run.clear()

    def resume(self):
        self._run.set()

    def cancel(self):
        self._cancel.set()
        self._run.set()

    @property
    def paused(self):
        return not self._run.is_set()

    def checkpoint(self):
        self._run.wait()
        if self._cancel.is_set():
            raise Cancelled()


@dataclass
class Progress:
    stage: str = ""
    book_index: int = 0
    book_total: int = 0
    book_name: str = ""
    book_pages_done: int = 0
    book_pages_total: int = 0
    pages_done: int = 0
    pages_total: int = 0
    ocr_pages: int = 0
    kruti_found: int = 0
    elapsed: float = 0.0
    eta: float | None = None


@dataclass
class RunResult:
    results_path: str
    krutis: int                # distinct Kruti (the same Kruti in several books counts once)
    places: int                # rows in the Excel: every book/page range where one was found
    books: int
    skipped: list = field(default_factory=list)
    in_excel: int = 0          # rows whose Kruti is in the input Excel


# ---------------------------------------------------------------- inputs
def _pdfs_in_folder(folder: Path) -> list:
    found = []
    for dirpath, _dirs, files in os.walk(folder, onerror=lambda e: None):  # skip unreadable
        found += [Path(dirpath) / f for f in files
                  if f.lower().endswith(".pdf") and not f.startswith("._")]
    return sorted(found)


def collect_pdfs(inputs) -> list:
    """Accept PDF files and/or folders (searched recursively). Duplicates removed."""
    seen, pdfs = set(), []
    if isinstance(inputs, (str, Path)):
        inputs = [inputs]
    for item in inputs:
        p = Path(item)
        found = _pdfs_in_folder(p) if p.is_dir() \
            else ([p] if p.suffix.lower() == ".pdf" and p.exists() else [])
        for x in found:
            key = str(x.resolve())
            if key not in seen:
                seen.add(key)
                pdfs.append(x)
    return pdfs


def load_krutis(excel_path, cfg):
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    sheet = cfg.get("excel_sheet", 0)
    ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not rows:
        raise ValueError("The Excel sheet is empty.")
    header = [str(h).strip() if h is not None else f"Column {i + 1}" for i, h in enumerate(rows[0])]
    wanted = [(cfg.get("excel_columns") or {}).get(k) for k in ("kruti_no", "name", "aadi", "ant")]
    idx = [header.index(w) if w in header else None for w in wanted]
    if None in idx:
        idx = [0, 1, 2, 3]            # fall back to the first four columns
    krutis = []
    for r in rows[1:]:
        r = (r + [None] * len(header))[:len(header)]
        aadi, ant = r[idx[2]], r[idx[3]]
        if not (aadi or ant):
            continue
        krutis.append({"id": r[idx[0]], "name": str(r[idx[1]] or "").strip(),
                       "aadi": str(aadi or ""), "ant": str(ant or ""),
                       "raw": ["" if v is None else v for v in r]})
    if not krutis:
        raise ValueError("No Kruti rows with Aadi/Ant Vakya found in the Excel.")
    return krutis, header


def book_name(pdf: Path) -> str:
    """Readable book name: the PDF's own title, else the file name without the library
    code and trailing catalogue number / scan-quality tags
    (B002637_abhay_ratna_sara_020001_hr3.pdf -> abhay ratna sara)."""
    try:
        import pymupdf
        with pymupdf.open(pdf) as doc:
            title = (doc.metadata or {}).get("title", "").strip()
        if title and not title.lower().endswith((".pdf", ".doc", ".docx")):
            return title
    except Exception:  # noqa: BLE001 - fall back to the file name
        pass
    name = re.sub(r"\s*\(\d+\)$", "", pdf.stem)                  # "file (1)" copies
    name = re.sub(r"^[A-Za-z]\d{4,}[_\- ]+", "", name)             # B002637_
    name = re.sub(r"_\d{5,}(_.*)?$", "", name)                     # _020001_hr3, _004503_<title>
    name = re.sub(r"_(\d+dpi|hr\d*|std|data|h_r|duplicate)$", "", name, flags=re.I)
    return re.sub(r"[_\-]+", " ", name).strip() or pdf.stem


# ---------------------------------------------------------------- matching
def _attach_matches(book, found, occurrences, krutis_by_id, page_script):
    """Mark each Kruti found in the book with the Excel Kruti whose start/end match covers
    it. Excel Kruti matched in the book but not split out as a block (no verse numbers)
    are added as rows of their own, so every match is listed."""
    used = set()
    for k in found:
        s, e = k.span
        best = None
        for i, (kid, o) in enumerate(occurrences):
            os_, oe = o.span_light
            overlap = max(0, min(e, oe) - max(s, os_))
            if overlap / max(1, e - s) >= 0.5 or s <= os_ < e:
                score = min(x for x in (o.aadi_score, o.ant_score) if x is not None)
                if best is None or score > best[0]:
                    best = (score, i, kid, o)
        if best:
            used.add(best[1])
            _set_match(k, krutis_by_id[best[2]], best[0])
    for i, (kid, o) in enumerate(occurrences):
        if i in used or o.status != FOUND:
            continue
        s, e = o.span_light
        lines = [ln.strip() for ln in book.light[s:e + 1].split("\n") if ln.strip()] or [""]
        k = Kruti(book.name, o.start_page, o.end_page, None, "", lines[0], lines[-1], "OK", (s, e))
        k.script = _script_for(k, page_script)
        _set_match(k, krutis_by_id[kid], min(o.aadi_score, o.ant_score))
        found.append(k)
    found.sort(key=lambda k: (k.start_page or 0, k.span[0]))


def _set_match(k, excel_kruti, score):
    k.kruti_no, k.kruti_name, k.match_score = excel_kruti["id"], excel_kruti["name"], score
    same = not excel_kruti["script"] or not k.script or excel_kruti["script"] == k.script
    k.match = "Matched - same script" if same else "Matched - other script"


def _script_for(k, page_script):
    pages = [page_script.get(p, "") for p in range(k.start_page or 0, (k.end_page or 0) + 1)]
    pages = [s for s in pages if s]
    return max(set(pages), key=pages.count) if pages else ""


# ---------------------------------------------------------------- run
def run(excel, inputs, out_dir, cfg, on_progress=None, log=print, ctl=None) -> RunResult:
    ctl = ctl or Controller()
    start = time.time()
    prog = Progress(stage="Reading Excel")
    last_emit = [0.0]
    sec_per_page = [None]      # moving average, for the ETA

    def emit(force=False):
        now = time.time()
        if on_progress and (force or now - last_emit[0] >= 0.15):
            prog.elapsed = now - start
            remaining = prog.pages_total - prog.pages_done
            prog.eta = remaining * sec_per_page[0] if sec_per_page[0] else None
            on_progress(prog)
            last_emit[0] = now

    emit(True)
    krutis, _columns = load_krutis(excel, cfg)
    for x in krutis:
        x["script"] = script_of(x["aadi"] + x["ant"])
    krutis_by_id = {x["id"]: x for x in krutis}
    log(f"{len(krutis)} Kruti read from {Path(excel).name}")

    pdfs = collect_pdfs(inputs)
    if not pdfs:
        raise ValueError("No PDF files found in the selected files/folders.")
    prog.stage = "Counting pages"
    emit(True)
    counts = []
    for p in pdfs:
        ctl.checkpoint()
        counts.append(page_count(str(p)))
    prog.book_total, prog.pages_total = len(pdfs), sum(counts)
    log(f"{len(pdfs)} books, {prog.pages_total} pages to scan")
    os.makedirs(out_dir, exist_ok=True)

    all_krutis, skipped = [], []
    pages_before = 0
    for n, (pdf, npages) in enumerate(zip(pdfs, counts), 1):
        ctl.checkpoint()
        name = pdf.stem
        prog.stage, prog.book_index, prog.book_name = "Reading pages", n, name
        prog.book_pages_done, prog.book_pages_total = 0, npages
        emit(True)
        t_last = [time.time()]

        def on_page(done, total, source, base=pages_before):
            now = time.time()
            if source != "cache":
                dt = now - t_last[0]
                sec_per_page[0] = dt if sec_per_page[0] is None else 0.9 * sec_per_page[0] + 0.1 * dt
            t_last[0] = now
            if source == "ocr":
                prog.ocr_pages += 1
            prog.book_pages_done, prog.book_pages_total = done, total
            prog.pages_done = base + done
            emit()

        try:
            pages = extract_book(str(pdf), cfg, on_page=on_page, checkpoint=ctl.checkpoint)
        except Cancelled:
            raise
        except Exception as e:  # noqa: BLE001 - a broken PDF must not stop a 200-book run
            log(f"[{n}/{len(pdfs)}] {name}: SKIPPED - could not read ({e})")
            skipped.append(name)
            pages_before += npages
            prog.pages_done = pages_before
            continue
        pages_before += npages
        prog.pages_done = pages_before

        prog.stage = "Finding Kruti"
        emit(True)
        book = Book.from_pages(pdf.name, pages)      # the file name is what the Excel shows
        page_script = {p.page: script_of(p.text) for p in pages}
        found = find_krutis(book, {p.page: p.confidence for p in pages}, cfg)
        for k in found:
            k.script = _script_for(k, page_script)
        index = BookIndex(book.key)
        occurrences = [(x["id"], o) for x in krutis for o in locate(x, book, index, cfg)
                       if o.span_light]
        _attach_matches(book, found, occurrences, krutis_by_id, page_script)
        title = book_name(pdf)
        for k in found:
            k.book_name = title
            if k.script == "Gujarati":            # show a Gujarati book's text in Gujarati
                k.title, k.aadi, k.ant = (devanagari_to_gujarati(t) for t in (k.title, k.aadi, k.ant))
        all_krutis += found
        prog.kruti_found = len(all_krutis)
        matched = sum(k.match != "New" for k in found)
        log(f"[{n}/{len(pdfs)}] {name}: {len(pages)} pages, {len(found)} Kruti "
            f"({matched} in the Excel, {len(found) - matched} new)")
        emit(True)

    prog.stage = "Writing Excel file"
    emit(True)
    all_krutis.sort(key=lambda k: (k.book, k.start_page or 0))
    assign_groups(all_krutis, cfg["duplicate_threshold"])
    results_path = os.path.join(out_dir, RESULTS_FILE)
    write_results(results_path, all_krutis)
    groups = {k.group_id for k in all_krutis}

    prog.stage = "Done"
    prog.pages_done = prog.pages_total
    emit(True)
    in_excel = sum(k.match != "New" for k in all_krutis)
    log(f"Done. {len(all_krutis)} Kruti found ({in_excel} in the Excel, "
        f"{len(all_krutis) - in_excel} new), {len(groups)} distinct.")
    return RunResult(results_path, len(groups), len(all_krutis), len(pdfs), skipped, in_excel)
